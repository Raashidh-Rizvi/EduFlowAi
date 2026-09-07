"""
===============================================================================
EduFlow AI - Slide & Document Parser (PDF & PowerPoint PPTX)
===============================================================================
Extracts structured text from lecture slide decks and documents:
- PDF files (.pdf) using pypdf
- PowerPoint presentations (.pptx) using python-pptx or native zip-xml parser
"""

import os
import re
import zipfile
import xml.etree.ElementTree as ET
from typing import List, Dict, Any, Optional


class SlidePage:
    def __init__(self, page_number: int, title: str, text: str, bullets: Optional[List[str]] = None):
        self.page_number = page_number
        self.title = title
        self.text = text
        self.bullets = bullets or []

    def to_dict(self) -> Dict[str, Any]:
        return {
            "page_number": self.page_number,
            "title": self.title,
            "text": self.text,
            "bullets": self.bullets
        }


class SlideParser:
    """
    Unified parser extracting slide-by-slide content from PDFs and PowerPoint presentations.
    """

    @staticmethod
    def extract_slides(file_path: str) -> List[SlidePage]:
        if not os.path.exists(file_path):
            raise FileNotFoundError(f"Slide file not found at path: {file_path}")

        ext = os.path.splitext(file_path)[1].lower()

        if ext == ".pdf":
            return SlideParser._extract_from_pdf(file_path)
        elif ext in [".pptx", ".ppt"]:
            return SlideParser._extract_from_pptx(file_path)
        else:
            # Try text or fallback
            return SlideParser._extract_generic_text(file_path)

    @staticmethod
    def _extract_from_pdf(file_path: str) -> List[SlidePage]:
        slides: List[SlidePage] = []
        try:
            from pypdf import PdfReader
            reader = PdfReader(file_path)
            for idx, page in enumerate(reader.pages):
                raw_text = page.extract_text() or ""
                lines = [l.strip() for l in raw_text.splitlines() if l.strip()]
                title = lines[0] if lines else f"Slide {idx + 1}"
                bullets = lines[1:] if len(lines) > 1 else []
                slides.append(SlidePage(
                    page_number=idx + 1,
                    title=title[:120],
                    text=raw_text,
                    bullets=bullets
                ))
        except Exception as e:
            print(f"[SlideParser] PDF extraction error: {e}")
        return slides

    @staticmethod
    def _extract_from_pptx(file_path: str) -> List[SlidePage]:
        slides: List[SlidePage] = []
        
        # 1. Try python-pptx if installed
        try:
            import pptx
            prs = pptx.Presentation(file_path)
            for idx, slide in enumerate(prs.slides):
                lines = []
                title = f"Slide {idx + 1}"
                for shape in slide.shapes:
                    if hasattr(shape, "text") and shape.text:
                        clean = shape.text.strip()
                        if clean:
                            lines.append(clean)
                if lines:
                    title = lines[0]
                    bullets = lines[1:]
                else:
                    bullets = []
                slides.append(SlidePage(
                    page_number=idx + 1,
                    title=title[:120],
                    text="\n".join(lines),
                    bullets=bullets
                ))
            if slides:
                return slides
        except ImportError:
            pass
        except Exception as e:
            print(f"[SlideParser] python-pptx error, falling back to zip-xml parser: {e}")

        # 2. Native ZIP/XML parsing for .pptx format (zero external dependencies)
        try:
            with zipfile.ZipFile(file_path, 'r') as z:
                slide_files = [f for f in z.namelist() if re.match(r'ppt/slides/slide\d+\.xml', f)]
                slide_files.sort(key=lambda x: int(re.search(r'\d+', x).group()))

                for idx, sfile in enumerate(slide_files):
                    content = z.read(sfile)
                    root = ET.fromstring(content)
                    text_elements = [elem.text for elem in root.iter() if elem.tag.endswith('}t') and elem.text]
                    clean_lines = [t.strip() for t in text_elements if t.strip()]
                    title = clean_lines[0] if clean_lines else f"Slide {idx + 1}"
                    bullets = clean_lines[1:] if len(clean_lines) > 1 else []
                    slides.append(SlidePage(
                        page_number=idx + 1,
                        title=title[:120],
                        text="\n".join(clean_lines),
                        bullets=bullets
                    ))
        except Exception as zip_e:
            print(f"[SlideParser] Zip XML PPTX extraction error: {zip_e}")

        return slides

    @staticmethod
    def _extract_generic_text(file_path: str) -> List[SlidePage]:
        try:
            with open(file_path, 'r', encoding='utf-8', errors='ignore') as f:
                content = f.read()
            paragraphs = [p.strip() for p in content.split("\n\n") if p.strip()]
            return [
                SlidePage(
                    page_number=i + 1,
                    title=f"Section {i + 1}",
                    text=p,
                    bullets=[]
                )
                for i, p in enumerate(paragraphs)
            ]
        except Exception as e:
            print(f"[SlideParser] Generic text extraction error: {e}")
            return []

    @staticmethod
    def get_full_text(slides: List[SlidePage]) -> str:
        parts = []
        for s in slides:
            parts.append(f"--- [Slide {s.page_number}: {s.title}] ---\n{s.text}")
        return "\n\n".join(parts)
